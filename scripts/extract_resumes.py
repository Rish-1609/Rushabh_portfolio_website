"""One-off helper: extract text from all resume PDFs/DOCX in the parent folder.

Writes a combined text dump to scripts/_extracted_resumes.txt for content authoring.
Not part of the website runtime.
"""
import os
import sys

from pypdf import PdfReader
import docx

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE_DIR = os.path.abspath(os.path.join(HERE, "..", ".."))
OUT = os.path.join(HERE, "_extracted_resumes.txt")


def read_pdf(path: str) -> str:
    try:
        reader = PdfReader(path)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:  # noqa: BLE001
        return f"[ERROR reading PDF: {exc}]"


def read_docx(path: str) -> str:
    try:
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:  # noqa: BLE001
        return f"[ERROR reading DOCX: {exc}]"


def main() -> None:
    files = sorted(os.listdir(SOURCE_DIR))
    chunks = []
    for name in files:
        if name.startswith("~"):
            continue
        path = os.path.join(SOURCE_DIR, name)
        if not os.path.isfile(path):
            continue
        lower = name.lower()
        if lower.endswith(".pdf"):
            text = read_pdf(path)
        elif lower.endswith(".docx"):
            text = read_docx(path)
        else:
            continue
        header = f"\n\n{'=' * 80}\nFILE: {name}\n{'=' * 80}\n"
        chunks.append(header + text.strip())

    with open(OUT, "w", encoding="utf-8") as fh:
        fh.write("\n".join(chunks))
    print(f"Wrote {OUT} ({len(chunks)} documents)")


if __name__ == "__main__":
    sys.exit(main())
